from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_new_feature(doc, text):
    """Add a paragraph with 'NEW:' prefix in green"""
    p = doc.add_paragraph()
    new_run = p.add_run('NEW: ')
    new_run.bold = True
    new_run.font.color.rgb = RGBColor(0, 128, 0)
    p.add_run(text)
    return p

def update_overview():
    doc = Document()

    # Title
    title = doc.add_heading('OneStat Platform Brief', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph('Overview Edition')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(14)
    subtitle.runs[0].font.color.rgb = RGBColor(128, 128, 128)

    doc.add_paragraph()

    # What It Is
    doc.add_heading('What It Is', 1)
    doc.add_paragraph(
        """OneStat is a full-stack GAA performance analytics platform built for county and club teams. """
        """It captures live match data, processes GPS and fitness information, runs agentic AI analysis """
        """across all data types, and delivers insights to coaches, analysts, and players through a single """
        """web application accessible on any device."""
    )

    # Core Capabilities
    doc.add_heading('Core Capabilities', 1)

    # Live Match Recording
    doc.add_heading('Live Match Recording', 2)
    doc.add_paragraph(
        """A touch-based pitch interface allows an analyst to record every event in real time — scores, """
        """wides, turnovers, kickouts, cards, fouls, substitutions, and possession. The ball is tracked """
        """across the pitch continuously. AI-generated insights are delivered every five minutes and at half time."""
    )
    doc.add_paragraph(
        """The pitch model is built to regulation GAA dimensions (145m × 88m) at a calibrated scale of 13.517px/m. """
        """Each tap stores an x/y coordinate that maps directly to a real-world location — at 7.4cm per pixel, """
        """every recorded event carries sub-metre positional accuracy of approximately ~15–20cm, enabling reliable """
        """spatial analysis in shot maps, heatmaps, and kickout distributions."""
    )

    # NEW features
    add_new_feature(doc,
        """Simple Scoring Mode — A streamlined phone-first recording interface for basic match coverage """
        """when full tactical tracking isn't needed. Tap-only scoring and events without continuous ball positioning."""
    )

    add_new_feature(doc,
        """Assist Tracking — Primary assists are explicitly captured and attributed during live recording, """
        """with dedicated leaderboards and season tracking."""
    )

    add_new_feature(doc,
        """Per-Match Pitch Dimensions — Optional real pitch measurements (130-145m × 80-90m) can be set """
        """per match for precise distance calculations on non-standard pitches."""
    )

    add_new_feature(doc, """Half-Time Features:""")
    doc.add_paragraph("""Screenshot sharing button — instantly capture and share half-time stats to coaching staff""", style='List Bullet')
    doc.add_paragraph("""Voice notes — record quick audio observations during the match via mic button""", style='List Bullet')
    doc.add_paragraph("""45m-line tap-to-place — explicit positioning for 45-metre kicks""", style='List Bullet')

    # Post-Match Analysis
    doc.add_heading('Post-Match Analysis', 2)
    doc.add_paragraph(
        """Following a match, the platform produces an AI-generated report covering scoring patterns, possession, """
        """kickout efficiency, turnover locations, and shot outcomes — all mapped visually on the pitch."""
    )

    add_new_feature(doc,
        """Player-Specific Charts — 6 dedicated player performance charts including shooting patterns, """
        """carry maps, and positional heatmaps."""
    )

    # GPS & Fitness Integration
    doc.add_heading('GPS & Fitness Integration', 2)
    doc.add_paragraph(
        """GPS files from wearable devices are uploaded and automatically processed. Player load, distance, """
        """speed zones, and heart rate data are correlated with match and training performance. Fitness test """
        """results are tracked longitudinally across the squad."""
    )

    # Season Intelligence
    doc.add_heading('Season Intelligence', 2)
    doc.add_paragraph(
        """A season-level AI dashboard surfaces KPI trends, outlier alerts, a weekly brief, and cross-data """
        """correlations (e.g. whether high training load predicts underperformance, or which players are in form)."""
    )

    add_new_feature(doc,
        """Orchestrator Leaderboard — New leaderboard category tracking playmakers by carry distance and progressive passes."""
    )

    # Video Analysis
    doc.add_heading('Video Analysis', 2)
    doc.add_paragraph(
        """Match video is uploaded and can be tagged manually or analysed automatically using AI computer vision — """
        """detecting player positions, classifying events, and generating a tactical bird's-eye view."""
    )

    add_new_feature(doc,
        """Video Compilations — AI-powered tool creates downloadable MP4 compilations from tagged clips """
        """(e.g., "show me all Conor Greene's wides this season"). Extraction uses HTTP range-seeking against """
        """R2 storage — never downloads full source videos. Cap: 20 clips per compilation."""
    )

    add_new_feature(doc, """Presentations — Full coaching presentation toolkit:""")
    doc.add_paragraph("""Slide deck builder with clip/animation/text slides""", style='List Bullet')
    doc.add_paragraph("""Freeze-frame annotation toolbar (arrows, lines, circles, spotlight, pencil, text)""", style='List Bullet')
    doc.add_paragraph("""Moving tracking graphics — keyframe-based player tracking rings""", style='List Bullet')
    doc.add_paragraph("""Per-slide voiceover recording""", style='List Bullet')
    doc.add_paragraph("""Player clip tagging with push notifications""", style='List Bullet')
    doc.add_paragraph('''"My Clips" tab in player portal for tagged clips''', style='List Bullet')

    # Set Piece & Tactical Preparation
    doc.add_heading('Set Piece & Tactical Preparation', 2)
    doc.add_paragraph(
        """Coaches build an animated playbook of set piece routines using a drag-and-drop pitch editor, """
        """record voiceovers, and push plays directly to players. An AI opposition briefing tool researches """
        """opponents using live web data."""
    )

    # Player Portal
    doc.add_heading('Player Portal', 2)
    doc.add_paragraph(
        """Players log in to a separate view showing their personal stats, leaderboard positions, and playbooks """
        """sent by the coaching staff."""
    )

    add_new_feature(doc,
        """Player Privacy Controls — Players can opt out of leaderboard visibility while still accessing """
        """their own stats and rankings."""
    )

    add_new_feature(doc,
        """My Clips Tab — Players see all presentation slides where they've been tagged by coaching staff."""
    )

    # AI Analyst Chat
    doc.add_heading('AI Analyst Chat', 2)
    doc.add_paragraph(
        """A conversational AI analyst — with full access to the club's match, training, GPS, and fitness data — """
        """answers natural language questions from coaches."""
    )

    # Technical Architecture
    doc.add_heading('Technical Architecture', 1)

    table = doc.add_table(rows=7, cols=2)
    table.style = 'Light Grid Accent 1'

    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Component'
    hdr_cells[1].text = 'Technology'

    rows_data = [
        ('Frontend', 'React (TypeScript), Vercel (global CDN)'),
        ('Backend API', 'Python / FastAPI, Fly.io (London region)'),
        ('Database', 'PostgreSQL, Supabase'),
        ('Object Storage', 'Cloudflare R2 (EU)'),
        ('Authentication', 'AWS Cognito — OAuth 2.0 with PKCE, httpOnly cookies'),
        ('AI', 'Anthropic Claude API (Claude Sonnet / Haiku)')
    ]

    for i, (comp, tech) in enumerate(rows_data, start=1):
        row_cells = table.rows[i].cells
        row_cells[0].text = comp
        row_cells[1].text = tech

    # Privacy & Compliance
    doc.add_heading('Privacy & Compliance', 1)

    add_new_feature(doc,
        """Opposition player data is automatically restricted to surname-only (enforced server-side) for GDPR compliance."""
    )

    add_new_feature(doc,
        """Weather logging now supports multi-select — multiple conditions (e.g., "windy + light rain") can be """
        """recorded per match."""
    )

    doc.add_paragraph()
    doc.add_paragraph()

    # Footer
    footer_p = doc.add_paragraph('OneStat — Turning GAA data into decisions.')
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.runs[0].italic = True

    conf_p = doc.add_paragraph('Confidential — Subject to NDA')
    conf_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    conf_p.runs[0].font.size = Pt(10)
    conf_p.runs[0].font.color.rgb = RGBColor(128, 128, 128)

    doc.save('C:/Users/owen_/Downloads/OneStat-Platform-Brief-Overview-UPDATED-Oct2026.docx')
    print('✓ Overview document updated and saved to Downloads')

if __name__ == '__main__':
    update_overview()
    print('\nDone! Check your Downloads folder for:')
    print('  - OneStat-Platform-Brief-Overview-UPDATED-Oct2026.docx')
