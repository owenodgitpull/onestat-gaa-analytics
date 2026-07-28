"""
Service layer for club onboarding operations.

Handles club creation, player file parsing (CSV/XLSX), and bulk player creation.
"""

import csv
import io
import logging
from datetime import date, datetime, timedelta
from typing import List, Optional, Tuple
from uuid import UUID, uuid4

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.club import Club
from app.models.player import Player, PlayerPosition, PlayerStatus
from app.schemas.onboarding import (
    ClubOnboardingCreate,
    PlayerPreviewRow,
    PlayerFilePreview,
    PlayerConfirmRow,
)

logger = logging.getLogger(__name__)

# Flexible column header matching (case-insensitive, stripped)
NAME_HEADERS = {"name", "player_name", "full_name", "player", "player name", "full name", "english name", "irish name"}
POSITION_HEADERS = {"position", "pos", "playing_position", "playing position"}
JERSEY_HEADERS = {"jersey_number", "number", "jersey", "no", "#", "jersey number", "no."}
DOB_HEADERS = {"date_of_birth", "dob", "birth_date", "birthday", "date of birth", "birth date", "d.o.b", "d.o.b.", "d.o.b / year", "d.o.b/ year"}
AGE_HEADERS = {"age", "player_age", "player age", "years", "age (approx)", "age(approx)", "approx age"}

# Position normalization
POSITION_ALIASES = {
    "gk": "goalkeeper", "keeper": "goalkeeper", "goalie": "goalkeeper", "1": "goalkeeper",
    "def": "defender", "back": "defender", "corner back": "defender", "cb": "defender",
    "half back": "defender", "hb": "defender", "full back": "defender", "fb": "defender",
    "wing back": "defender", "wb": "defender",
    "mid": "midfielder", "midfield": "midfielder", "centre": "midfielder",
    "half forward": "midfielder", "hf": "midfielder",
    "fwd": "forward", "forward": "forward", "corner forward": "forward", "cf": "forward",
    "full forward": "forward", "ff": "forward", "wing forward": "forward", "wf": "forward",
}


class OnboardingService:
    """Service for club onboarding operations."""

    @staticmethod
    async def create_club(db: AsyncSession, data: ClubOnboardingCreate) -> Club:
        """Create a new club during onboarding."""
        club = Club(
            name=data.name,
            short_name=data.short_name,
            county=data.county,
            province=data.province,
            home_ground=data.home_ground,
            primary_colour=data.primary_colour,
            secondary_colour=data.secondary_colour,
            is_active=True,
            onboarding_completed=False,
            trial_ends_at=datetime.utcnow() + timedelta(days=30),
        )
        db.add(club)
        await db.commit()
        await db.refresh(club)
        return club

    @staticmethod
    async def update_club_logo(db: AsyncSession, club_id: UUID, logo_url: str) -> Club:
        """Update club logo URL."""
        result = await db.execute(select(Club).where(Club.id == club_id))
        club = result.scalar_one_or_none()
        if not club:
            raise ValueError(f"Club {club_id} not found")
        club.logo_url = logo_url
        await db.commit()
        await db.refresh(club)
        return club

    @staticmethod
    def parse_player_file(file_content: bytes, filename: str, name_column: Optional[str] = None) -> PlayerFilePreview:
        """
        Parse a CSV or XLSX file into a player preview.

        Supports flexible column headers and normalizes positions.
        Returns warnings for invalid/ambiguous data.
        If name_column is provided, use that header as the name field.
        """
        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

        if ext in ("xlsx", "xls"):
            return OnboardingService._parse_xlsx(file_content, name_column=name_column)
        else:
            return OnboardingService._parse_csv(file_content, name_column=name_column)

    @staticmethod
    def _parse_csv(content: bytes, name_column: Optional[str] = None) -> PlayerFilePreview:
        """Parse CSV content into player preview rows."""
        text = content.decode("utf-8-sig")  # Handle BOM
        # Detect delimiter
        sniffer = csv.Sniffer()
        try:
            dialect = sniffer.sniff(text[:2048])
        except csv.Error:
            dialect = csv.excel

        reader = csv.reader(io.StringIO(text), dialect)
        rows_raw = list(reader)

        if len(rows_raw) < 2:
            return PlayerFilePreview(parsed_count=0, valid_count=0, warnings=["File is empty or has only headers"], rows=[])

        # Auto-detect header row (may not be the first row)
        header_idx = OnboardingService._find_header_row(rows_raw)
        headers = [h.strip().lower() for h in rows_raw[header_idx]]
        original_headers = [h.strip() for h in rows_raw[header_idx]]
        col_map = OnboardingService._map_columns(headers)

        # Apply user override for name column
        if name_column:
            nc = name_column.strip().lower()
            for idx, h in enumerate(headers):
                if h == nc:
                    col_map["name"] = idx
                    break

        if "name" not in col_map and "_split_name" not in col_map:
            return PlayerFilePreview(
                parsed_count=0, valid_count=0,
                warnings=["Could not auto-detect name column. Please select which column contains player names."],
                rows=[],
                headers=original_headers,
                column_mapping={k: v for k, v in col_map.items() if not k.startswith("_")},
                needs_name_column=True,
            )

        result = OnboardingService._process_rows(rows_raw[header_idx + 1:], col_map)
        result.headers = original_headers
        result.column_mapping = {k: v for k, v in col_map.items() if not k.startswith("_")}
        return result

    @staticmethod
    def _parse_xlsx(content: bytes, name_column: Optional[str] = None) -> PlayerFilePreview:
        """Parse XLSX content into player preview rows."""
        try:
            from openpyxl import load_workbook
        except ImportError:
            return PlayerFilePreview(
                parsed_count=0, valid_count=0,
                warnings=["openpyxl is required for XLSX parsing. Install with: pip install openpyxl"],
                rows=[],
            )

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        if ws is None:
            return PlayerFilePreview(parsed_count=0, valid_count=0, warnings=["No active worksheet found"], rows=[])

        rows_raw = []
        for row in ws.iter_rows(values_only=True):
            rows_raw.append([str(cell) if cell is not None else "" for cell in row])

        wb.close()

        if len(rows_raw) < 2:
            return PlayerFilePreview(parsed_count=0, valid_count=0, warnings=["File is empty or has only headers"], rows=[])

        # Auto-detect header row (may not be the first row)
        header_idx = OnboardingService._find_header_row(rows_raw)
        headers = [h.strip().lower() for h in rows_raw[header_idx]]
        original_headers = [h.strip() for h in rows_raw[header_idx]]
        col_map = OnboardingService._map_columns(headers)

        # Apply user override for name column
        if name_column:
            nc = name_column.strip().lower()
            for idx, h in enumerate(headers):
                if h == nc:
                    col_map["name"] = idx
                    break

        if "name" not in col_map and "_split_name" not in col_map:
            return PlayerFilePreview(
                parsed_count=0, valid_count=0,
                warnings=["Could not auto-detect name column. Please select which column contains player names."],
                rows=[],
                headers=original_headers,
                column_mapping={k: v for k, v in col_map.items() if not k.startswith("_")},
                needs_name_column=True,
            )

        result = OnboardingService._process_rows(rows_raw[header_idx + 1:], col_map)
        result.headers = original_headers
        result.column_mapping = {k: v for k, v in col_map.items() if not k.startswith("_")}
        return result

    @staticmethod
    def _find_header_row(rows_raw: List[List[str]], max_scan: int = 10) -> int:
        """
        Scan the first few rows to find the actual header row.
        Returns the 0-based index of the row containing recognizable column headers.
        Falls back to 0 if no header row is found.
        """
        all_known = NAME_HEADERS | POSITION_HEADERS | JERSEY_HEADERS | DOB_HEADERS | AGE_HEADERS
        all_known |= {"first name", "first_name", "firstname", "forename",
                       "last name", "last_name", "lastname", "surname"}
        for i, row in enumerate(rows_raw[:max_scan]):
            normalized = [cell.strip().lower() for cell in row]
            matches = sum(1 for cell in normalized if cell in all_known)
            if matches >= 2:  # At least 2 recognized headers
                return i
        return 0

    @staticmethod
    def _map_columns(headers: List[str]) -> dict:
        """Map header names to column indices."""
        FIRST_NAME_HEADERS = {"first name", "first_name", "firstname", "forename"}
        LAST_NAME_HEADERS = {"last name", "last_name", "lastname", "surname"}
        col_map = {}
        for idx, h in enumerate(headers):
            if h in NAME_HEADERS:
                col_map["name"] = idx
            elif h in FIRST_NAME_HEADERS:
                col_map["first_name"] = idx
            elif h in LAST_NAME_HEADERS:
                col_map["last_name"] = idx
            elif h in POSITION_HEADERS:
                col_map["position"] = idx
            elif h in JERSEY_HEADERS:
                col_map["jersey"] = idx
            elif h in DOB_HEADERS:
                col_map["dob"] = idx
            elif h in AGE_HEADERS:
                col_map["age"] = idx
        # Synthesize "name" from first + last if no single name column
        if "name" not in col_map and "first_name" in col_map:
            col_map["_split_name"] = True
        return col_map

    @staticmethod
    def _process_rows(data_rows: List[List[str]], col_map: dict) -> PlayerFilePreview:
        """Process data rows into player preview rows with warnings."""
        preview_rows = []
        file_warnings = []
        valid_count = 0

        for i, row in enumerate(data_rows):
            row_num = i + 2  # 1-indexed, +1 for header
            warnings = []

            # Name (required) — single column or first+last
            if "_split_name" in col_map:
                first = row[col_map["first_name"]].strip() if col_map["first_name"] < len(row) else ""
                last = row[col_map.get("last_name", -1)].strip() if col_map.get("last_name", -1) < len(row) else ""
                name = f"{first} {last}".strip()
            else:
                name = row[col_map["name"]].strip() if col_map["name"] < len(row) else ""
            if not name:
                continue  # Skip blank rows

            # Position
            position = None
            if "position" in col_map and col_map["position"] < len(row):
                raw_pos = row[col_map["position"]].strip().lower()
                if raw_pos:
                    position = POSITION_ALIASES.get(raw_pos, raw_pos)
                    valid_positions = {p.value for p in PlayerPosition}
                    if position not in valid_positions:
                        warnings.append(f"Unknown position '{row[col_map['position']].strip()}' — will default to null")
                        position = None

            # Jersey number
            jersey_number = None
            if "jersey" in col_map and col_map["jersey"] < len(row):
                raw_jersey = row[col_map["jersey"]].strip()
                if raw_jersey:
                    try:
                        jersey_number = int(float(raw_jersey))
                        if jersey_number < 1 or jersey_number > 99:
                            warnings.append(f"Jersey number {jersey_number} out of range (1-99)")
                            jersey_number = None
                    except (ValueError, TypeError):
                        warnings.append(f"Invalid jersey number '{raw_jersey}'")

            # DOB (prefer explicit DOB over age)
            dob_str = None
            if "dob" in col_map and col_map["dob"] < len(row):
                raw_dob = row[col_map["dob"]].strip()
                if raw_dob and raw_dob not in ("—", "-", "–", "N/A", "n/a", "None", "none", ""):
                    # Skip warning for clearly non-date values (e.g. ~2003)
                    dob_str = OnboardingService._parse_date(raw_dob)
                    if dob_str is None and not raw_dob.startswith("~"):
                        warnings.append(f"Unparseable date '{raw_dob}'")

            # Age → approximate DOB (Jan 1 of birth year) if no explicit DOB
            if dob_str is None and "age" in col_map and col_map["age"] < len(row):
                raw_age = row[col_map["age"]].strip()
                if raw_age and raw_age not in ("—", "-", "–", "N/A", "n/a", "None", "none"):
                    try:
                        age_val = int(float(raw_age))
                        if 10 <= age_val <= 60:
                            birth_year = date.today().year - age_val
                            dob_str = f"{birth_year}-01-01"
                        else:
                            warnings.append(f"Age {age_val} out of reasonable range (10-60)")
                    except (ValueError, TypeError):
                        pass  # Silently skip non-numeric age values

            preview_rows.append(PlayerPreviewRow(
                row_number=row_num,
                name=name,
                position=position,
                jersey_number=jersey_number,
                date_of_birth=dob_str,
                warnings=warnings,
            ))

            if not warnings:
                valid_count += 1

        return PlayerFilePreview(
            parsed_count=len(preview_rows),
            valid_count=valid_count,
            warnings=file_warnings,
            rows=preview_rows,
        )

    @staticmethod
    def _parse_date(raw: str) -> Optional[str]:
        """Try to parse a date string into YYYY-MM-DD format."""
        formats = [
            "%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y",
            "%d %b %Y", "%d %B %Y", "%Y/%m/%d", "%d.%m.%Y",
        ]
        for fmt in formats:
            try:
                d = datetime.strptime(raw, fmt).date()
                return d.isoformat()
            except ValueError:
                continue
        return None

    @staticmethod
    async def bulk_create_players(
        db: AsyncSession,
        club_id: UUID,
        players: List[PlayerConfirmRow],
    ) -> List[Player]:
        """Bulk create players for a club."""
        created = []
        for p in players:
            dob = None
            if p.date_of_birth:
                try:
                    dob = date.fromisoformat(p.date_of_birth)
                except ValueError:
                    pass

            kwargs: dict = dict(
                id=uuid4(),
                club_id=club_id,
                name=p.name,
                jersey_number=p.jersey_number,
                date_of_birth=dob,
                status=PlayerStatus.ACTIVE,
                active=True,
            )
            if p.position:
                kwargs["position"] = p.position
            player = Player(**kwargs)
            db.add(player)
            created.append(player)

        await db.commit()
        for player in created:
            await db.refresh(player)

        logger.info(f"Bulk created {len(created)} players for club {club_id}")
        return created

    @staticmethod
    async def bulk_upsert_players(
        db: AsyncSession,
        club_id: UUID,
        players: List[PlayerConfirmRow],
    ) -> dict:
        """
        Upsert players for a club.
        - Match existing players by normalized name (case-insensitive, stripped).
        - Update position/jersey/dob on existing players if the new file provides them.
        - Create new players that don't exist yet.
        - Never delete or deactivate existing players.
        Returns: { created: int, updated: int, unchanged: int, details: [...] }
        """
        # Fetch all existing players for this club
        result = await db.execute(
            select(Player).where(Player.club_id == club_id)
        )
        existing = list(result.scalars().all())
        existing_by_name: dict[str, Player] = {}
        for p in existing:
            existing_by_name[p.name.strip().lower()] = p

        created, updated, unchanged = 0, 0, 0
        details = []

        for row in players:
            norm_name = row.name.strip().lower()
            dob = None
            if row.date_of_birth:
                try:
                    dob = date.fromisoformat(row.date_of_birth)
                except ValueError:
                    pass

            if norm_name in existing_by_name:
                # Existing player — update fields if file provides new data
                player = existing_by_name[norm_name]
                changed = False
                if row.position and player.position != row.position:
                    player.position = row.position
                    changed = True
                if row.jersey_number is not None and player.jersey_number != row.jersey_number:
                    player.jersey_number = row.jersey_number
                    changed = True
                if dob and player.date_of_birth != dob:
                    player.date_of_birth = dob
                    changed = True
                # Reactivate if was soft-deleted
                if not player.active:
                    player.active = True
                    player.status = PlayerStatus.ACTIVE
                    changed = True

                if changed:
                    updated += 1
                    details.append({"name": player.name, "action": "updated"})
                else:
                    unchanged += 1
                    details.append({"name": player.name, "action": "unchanged"})
            else:
                # New player — only set position if provided (DB column is enum, can't cast None via VARCHAR)
                kwargs: dict = dict(
                    id=uuid4(),
                    club_id=club_id,
                    name=row.name.strip(),
                    jersey_number=row.jersey_number,
                    date_of_birth=dob,
                    status=PlayerStatus.ACTIVE,
                    active=True,
                )
                if row.position:
                    kwargs["position"] = row.position
                player = Player(**kwargs)
                db.add(player)
                created += 1
                details.append({"name": player.name, "action": "created"})

        await db.commit()
        logger.info(
            f"Roster upsert for club {club_id}: {created} created, {updated} updated, {unchanged} unchanged"
        )
        return {"created": created, "updated": updated, "unchanged": unchanged, "details": details}

    @staticmethod
    async def complete_onboarding(db: AsyncSession, club_id: UUID) -> Club:
        """Mark club onboarding as completed."""
        result = await db.execute(select(Club).where(Club.id == club_id))
        club = result.scalar_one_or_none()
        if not club:
            raise ValueError(f"Club {club_id} not found")
        club.onboarding_completed = True
        await db.commit()
        await db.refresh(club)
        return club
