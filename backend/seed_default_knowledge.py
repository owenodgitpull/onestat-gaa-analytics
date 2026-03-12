"""
One-time script to seed default knowledge base documents from R2.

Reads PDFs from the R2 folder 'default_gaa_football_knowledge_base/',
creates KnowledgeDocument records (is_default=True, club_id=NULL),
and processes them into RAG chunks accessible to all teams.

Usage:
    cd backend
    python seed_default_knowledge.py
"""

import asyncio
import logging
import sys

logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
logger = logging.getLogger(__name__)

# Ensure app modules are importable
sys.path.insert(0, ".")

R2_PREFIX = "default_gaa_football_knowledge_base/"


def guess_doc_type(filename: str) -> str:
    name = filename.lower()
    if "rule" in name:
        return "rules"
    if "tactic" in name or "coaching" in name or "blueprint" in name:
        return "tactics"
    if "gps" in name or "statsport" in name or "fitness" in name:
        return "statsports"
    return "other"


async def main():
    from app.database import async_session_maker
    from app.models.knowledge_document import KnowledgeDocument
    from app.services.storage_service import storage
    from app.services.rag_service import RAGService
    from sqlalchemy import select

    if not storage.is_configured:
        logger.error("R2 storage not configured. Check your .env for R2 credentials.")
        return

    # List files in defaults folder
    files = storage.list_files(prefix=R2_PREFIX)
    if not files:
        logger.error(f"No files found at R2 prefix: {R2_PREFIX}")
        return

    logger.info(f"Found {len(files)} files in {R2_PREFIX}")

    import fitz  # pymupdf
    import io

    for f in files:
        r2_key = f["key"]
        filename = r2_key.split("/")[-1]
        if not filename:
            continue

        # Skip empty folder markers and non-documents
        if f["size"] == 0:
            continue
        is_pdf = filename.lower().endswith(".pdf")
        is_docx = filename.lower().endswith(".docx")
        if not is_pdf and not is_docx:
            logger.info(f"  Skipping non-document: {filename}")
            continue

        # Each doc gets its own session to avoid transaction contamination
        async with async_session_maker() as db:
            try:
                # Check if already seeded
                existing = await db.execute(
                    select(KnowledgeDocument).where(
                        KnowledgeDocument.is_default.is_(True),
                        KnowledgeDocument.original_filename == filename,
                    )
                )
                if existing.scalar_one_or_none():
                    logger.info(f"  Already seeded: {filename} — skipping")
                    continue

                logger.info(f"  Processing: {filename} ({f['size']} bytes)")

                # Download from R2
                file_bytes = storage.download_file(r2_key)
                if not file_bytes:
                    logger.error(f"  Failed to download {r2_key}")
                    continue

                # Extract text
                text_content = ""
                if is_pdf:
                    text_parts = []
                    with fitz.open(stream=io.BytesIO(file_bytes), filetype="pdf") as pdf_doc:
                        for page_num, page in enumerate(pdf_doc, 1):
                            page_text = page.get_text()
                            if page_text.strip():
                                text_parts.append(f"[Page {page_num}]\n{page_text}")
                    text_content = "\n\n".join(text_parts)
                elif is_docx:
                    import docx
                    doc_file = docx.Document(io.BytesIO(file_bytes))
                    paragraphs = [p.text for p in doc_file.paragraphs if p.text.strip()]
                    text_content = "\n\n".join(paragraphs)

                if not text_content or len(text_content.strip()) < 50:
                    logger.error(f"  Could not extract meaningful text from {filename}")
                    continue

                logger.info(f"  Extracted {len(text_content)} chars of text")

                content_type = "application/pdf" if is_pdf else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                doc_type = guess_doc_type(filename)

                # Create KnowledgeDocument record
                doc = KnowledgeDocument(
                    club_id=None,
                    filename=filename,
                    original_filename=filename,
                    doc_type=doc_type,
                    r2_key=r2_key,
                    content_type=content_type,
                    is_default=True,
                    processing_status="processing",
                    file_size_bytes=len(file_bytes),
                )
                db.add(doc)
                await db.commit()
                await db.refresh(doc)

                # Process through RAG — club_id=None for shared chunks
                rag_result = await RAGService.process_document(
                    db,
                    source_file=filename,
                    content=text_content,
                    doc_type=doc_type,
                    force=True,
                    club_id=None,
                )

                doc.chunk_count = rag_result.get("chunk_count", 0)
                doc.processing_status = "completed"
                await db.commit()

                logger.info(f"  Done: {filename} — {doc.chunk_count} chunks, type={doc_type}")

            except Exception as e:
                logger.error(f"  FAILED: {filename} — {e}")
                await db.rollback()

    logger.info("Seed complete!")


if __name__ == "__main__":
    asyncio.run(main())
